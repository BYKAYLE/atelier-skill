#!/usr/bin/env python3
"""
build_slide.py - python-pptx 기반 슬라이드 디자인 라이브러리 + 빌더
Usage: python3 build_slide.py --plan <plan.json> --output <output.pptx>
"""

import argparse
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# 기본 색상 팔레트 (다크 테마)
# ---------------------------------------------------------------------------

DEFAULT_COLORS = {
    "bg": "0B0F1A",
    "bg_card": "121A2D",
    "bg_card_light": "1A243B",
    "text": "FFFFFF",
    "text_secondary": "B0B8C8",
    "text_muted": "6B7280",
    "accent1": "4A90FF",
    "accent2": "00D4FF",
    "accent3": "7C5CFC",
    "accent4": "FFC107",
    "accent5": "00E676",
    "accent6": "FF5252",
    "card_border": "253048",
}

# ---------------------------------------------------------------------------
# ColorPalette
# ---------------------------------------------------------------------------

class ColorPalette:
    """slide_plan의 global_colors에서 초기화"""

    def __init__(self, colors_dict, font_name="Apple SD Gothic Neo"):
        self.font = font_name
        merged = {**DEFAULT_COLORS, **colors_dict}
        for key, val in merged.items():
            setattr(self, key, RGBColor.from_string(val))

    def get(self, key):
        """키 이름으로 RGBColor 반환. 없으면 흰색."""
        return getattr(self, key, RGBColor(0xFF, 0xFF, 0xFF))


# ---------------------------------------------------------------------------
# 유틸리티 함수
# ---------------------------------------------------------------------------

def set_slide_bg(slide, color):
    """슬라이드 배경색 설정"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_glow_circle(slide, left, top, size, color, alpha=0.15):
    """반투명 글로우 원 추가"""
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE -> use freeform circle via oval
        left, top, size, size,
    )
    # oval via XML adjustment
    sp = shape._element
    prstGeom = sp.spPr.find(qn("a:prstGeom"))
    if prstGeom is not None:
        prstGeom.set("prst", "ellipse")

    spPr = shape._element.spPr
    # remove existing solidFill
    existing = spPr.find(qn("a:solidFill"))
    if existing is not None:
        spPr.remove(existing)

    solidFill = etree.SubElement(spPr, qn("a:solidFill"))
    srgb = etree.SubElement(solidFill, qn("a:srgbClr"))
    srgb.set("val", f"{color.rgb:06X}" if hasattr(color, "rgb") else f"{int(color[0]):02X}{int(color[1]):02X}{int(color[2]):02X}")
    alpha_elem = etree.SubElement(srgb, qn("a:alpha"))
    alpha_elem.set("val", str(int(alpha * 100000)))

    # no line
    ln = spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(spPr, qn("a:ln"))
    noFill = etree.SubElement(ln, qn("a:noFill"))
    return shape


def _rgb_hex(color):
    """RGBColor -> 6자리 hex 문자열"""
    try:
        return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
    except Exception:
        return "FFFFFF"


def add_glow_circle_v2(slide, left, top, size, color, alpha=0.15):
    """반투명 글로우 원 추가 (RGBColor 객체 지원)"""
    from pptx.util import Emu
    shape = slide.shapes.add_shape(
        1, left, top, size, size,
    )
    sp = shape._element
    prstGeom = sp.spPr.find(qn("a:prstGeom"))
    if prstGeom is not None:
        prstGeom.set("prst", "ellipse")

    spPr = shape._element.spPr
    existing = spPr.find(qn("a:solidFill"))
    if existing is not None:
        spPr.remove(existing)

    hex_val = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
    solidFill = etree.SubElement(spPr, qn("a:solidFill"))
    srgb = etree.SubElement(solidFill, qn("a:srgbClr"))
    srgb.set("val", hex_val)
    alpha_elem = etree.SubElement(srgb, qn("a:alpha"))
    alpha_elem.set("val", str(int(alpha * 100000)))

    ln = spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(spPr, qn("a:ln"))
    noFill = ln.find(qn("a:noFill"))
    if noFill is None:
        etree.SubElement(ln, qn("a:noFill"))

    shape.line.fill.background()
    return shape


def add_accent_line(slide, left, top, width, color, height=None):
    """강조 수평선 추가"""
    if height is None:
        height = Pt(3)
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_card(slide, left, top, width, height, color, border_color=None):
    """카드 사각형 추가"""
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    # 라운드 코너
    sp = shape._element
    prstGeom = sp.spPr.find(qn("a:prstGeom"))
    if prstGeom is not None:
        prstGeom.set("prst", "roundRect")
        avLst = prstGeom.find(qn("a:avLst"))
        if avLst is None:
            avLst = etree.SubElement(prstGeom, qn("a:avLst"))
        gd = avLst.find(qn("a:gd"))
        if gd is None:
            gd = etree.SubElement(avLst, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", "val 5000")
    return shape


def add_text_box(slide, left, top, width, height):
    """텍스트 박스 추가"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    return txBox


def set_text(tf, text, size=18, bold=False, color=None, align=PP_ALIGN.LEFT, font_name=None):
    """텍스트 프레임의 첫번째 단락 텍스트 설정"""
    if color is None:
        color = RGBColor(0xFF, 0xFF, 0xFF)
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = align
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if font_name:
            run.font.name = font_name


def add_paragraph(tf, text, size=14, bold=False, color=None, align=PP_ALIGN.LEFT,
                  space_before=0, space_after=0, font_name=None):
    """텍스트 프레임에 단락 추가"""
    if color is None:
        color = RGBColor(0xFF, 0xFF, 0xFF)
    p = tf.add_paragraph()
    p.text = text
    p.alignment = align
    if space_before:
        p.space_before = Pt(space_before)
    if space_after:
        p.space_after = Pt(space_after)
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if font_name:
            run.font.name = font_name
    return p


def add_badge(slide, left, top, text, color, text_color=None, width=None):
    """뱃지 (라운드 사각형) 추가"""
    if text_color is None:
        text_color = RGBColor(0xFF, 0xFF, 0xFF)
    badge_h = Pt(22)
    if width is None:
        width = Inches(2.2)
    shape = slide.shapes.add_shape(1, left, top, width, badge_h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp = shape._element
    prstGeom = sp.spPr.find(qn("a:prstGeom"))
    if prstGeom is not None:
        prstGeom.set("prst", "roundRect")
        avLst = prstGeom.find(qn("a:avLst"))
        if avLst is None:
            avLst = etree.SubElement(prstGeom, qn("a:avLst"))
        gd = avLst.find(qn("a:gd"))
        if gd is None:
            gd = etree.SubElement(avLst, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", "val 30000")
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = text_color
    if True:
        pass
    return shape


def add_stat_number(slide, left, top, number, label, color, width):
    """큰 숫자 + 라벨 추가"""
    tb = add_text_box(slide, left, top, width, Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = number
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = color

    tb2 = add_text_box(slide, left, top + Inches(0.65), width, Inches(0.35))
    tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]
    run2 = p2.add_run()
    run2.text = label
    run2.font.size = Pt(12)
    run2.font.color.rgb = RGBColor(0xB0, 0xB8, 0xC8)
    return tb, tb2


def add_progress_bar(slide, left, top, width, fill_ratio, color, height=None):
    """프로그레스 바 추가 (배경 + 채움)"""
    if height is None:
        height = Pt(8)
    bg = slide.shapes.add_shape(1, left, top, width, height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0x25, 0x30, 0x48)
    bg.line.fill.background()
    sp = bg._element
    prstGeom = sp.spPr.find(qn("a:prstGeom"))
    if prstGeom is not None:
        prstGeom.set("prst", "roundRect")

    fill_width = max(int(width * fill_ratio), Pt(4))
    fg = slide.shapes.add_shape(1, left, top, fill_width, height)
    fg.fill.solid()
    fg.fill.fore_color.rgb = color
    fg.line.fill.background()
    sp2 = fg._element
    prstGeom2 = sp2.spPr.find(qn("a:prstGeom"))
    if prstGeom2 is not None:
        prstGeom2.set("prst", "roundRect")
    return bg, fg


def add_icon_circle(slide, left, top, icon_text, color, size=None):
    """아이콘 원 추가"""
    if size is None:
        size = Inches(0.7)
    shape = slide.shapes.add_shape(1, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp = shape._element
    prstGeom = sp.spPr.find(qn("a:prstGeom"))
    if prstGeom is not None:
        prstGeom.set("prst", "ellipse")
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = icon_text
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return shape


def add_notes(slide, text):
    """발표자 노트 추가"""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = text


def _resolve_color(key_or_hex, palette):
    """accent1 같은 키 또는 HEX 문자열을 RGBColor로 변환"""
    if key_or_hex and hasattr(palette, key_or_hex):
        return palette.get(key_or_hex)
    try:
        return RGBColor.from_string(key_or_hex)
    except Exception:
        return palette.get("accent1")


def _add_title_section(slide, spec, palette, y_start=None):
    """공통 뱃지 + 타이틀 + 서브타이틀 섹션"""
    if y_start is None:
        y_start = Inches(0.5)
    margin_l = Inches(0.8)
    slide_w = Inches(13.333)
    content_w = slide_w - margin_l * 2

    badge_text = spec.get("badge", "")
    badge_color_key = spec.get("badge_color", "accent1")
    badge_color = _resolve_color(badge_color_key, palette)

    y = y_start
    if badge_text:
        add_badge(slide, margin_l, y, badge_text, badge_color, RGBColor(0xFF, 0xFF, 0xFF))
        y += Inches(0.42)

    title = spec.get("title", "")
    if title:
        tb = add_text_box(slide, margin_l, y, content_w, Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = title
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = palette.get("text")
        run.font.name = palette.font
        y += Inches(0.65)

    subtitle = spec.get("subtitle", "")
    if subtitle:
        tb2 = add_text_box(slide, margin_l, y, content_w, Inches(0.45))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.LEFT
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(14)
        run2.font.color.rgb = palette.get("text_secondary")
        run2.font.name = palette.font
        y += Inches(0.45)

    # 프로급: 타이틀 섹션 하단 그라데이션 강조선
    if title:
        add_gradient_line(slide, margin_l, y, Inches(4),
                          palette.get("accent1"), palette.get("accent2"), Pt(2))
        y += Inches(0.2)

    return y


# ---------------------------------------------------------------------------
# 프로급 디자인 유틸리티 (v2.0 강화)
# ---------------------------------------------------------------------------

def set_gradient_bg(slide, color1, color2, angle=5400000):
    """슬라이드 그라데이션 배경 (linear). angle: 60000ths of degree, 5400000=90도(위→아래)"""
    bg = slide.background
    bgPr = bg._element
    # Remove existing bgPr/bgFillStyleLst if any
    for child in list(bgPr):
        bgPr.remove(child)
    # Build gradFill
    bgPr_elem = etree.SubElement(bgPr, qn("p:bgPr"))
    gradFill = etree.SubElement(bgPr_elem, qn("a:gradFill"))
    gradFill.set("flip", "none")
    gradFill.set("rotWithShape", "0")
    gsLst = etree.SubElement(gradFill, qn("a:gsLst"))
    # Stop 1
    gs1 = etree.SubElement(gsLst, qn("a:gs"))
    gs1.set("pos", "0")
    srgb1 = etree.SubElement(gs1, qn("a:srgbClr"))
    srgb1.set("val", _rgb_hex(color1))
    # Stop 2
    gs2 = etree.SubElement(gsLst, qn("a:gs"))
    gs2.set("pos", "100000")
    srgb2 = etree.SubElement(gs2, qn("a:srgbClr"))
    srgb2.set("val", _rgb_hex(color2))
    # Linear direction
    lin = etree.SubElement(gradFill, qn("a:lin"))
    lin.set("ang", str(angle))
    lin.set("scaled", "1")
    # Required a:effectLst
    etree.SubElement(bgPr_elem, qn("a:effectLst"))


def add_card_shadow(shape, blur=50800, dist=38100, direction=2700000, alpha=0.4):
    """카드에 drop shadow 효과 추가. blur/dist in EMU."""
    sp = shape._element
    spPr = sp.spPr
    effectLst = spPr.find(qn("a:effectLst"))
    if effectLst is None:
        effectLst = etree.SubElement(spPr, qn("a:effectLst"))
    outerShdw = etree.SubElement(effectLst, qn("a:outerShdw"))
    outerShdw.set("blurRad", str(blur))
    outerShdw.set("dist", str(dist))
    outerShdw.set("dir", str(direction))
    outerShdw.set("algn", "tl")
    outerShdw.set("rotWithShape", "0")
    srgb = etree.SubElement(outerShdw, qn("a:srgbClr"))
    srgb.set("val", "000000")
    alpha_elem = etree.SubElement(srgb, qn("a:alpha"))
    alpha_elem.set("val", str(int(alpha * 100000)))


def add_glass_card(slide, left, top, width, height, color, alpha=0.25, border_color=None, border_alpha=0.3):
    """글래스모피즘 카드 (반투명 + 미묘한 보더 + 라운드 + 그림자)"""
    shape = slide.shapes.add_shape(1, left, top, width, height)
    sp = shape._element
    spPr = sp.spPr

    # roundRect
    prstGeom = spPr.find(qn("a:prstGeom"))
    if prstGeom is not None:
        prstGeom.set("prst", "roundRect")
        avLst = prstGeom.find(qn("a:avLst"))
        if avLst is None:
            avLst = etree.SubElement(prstGeom, qn("a:avLst"))
        gd = avLst.find(qn("a:gd"))
        if gd is None:
            gd = etree.SubElement(avLst, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", "val 5000")

    # semi-transparent fill
    existing_fill = spPr.find(qn("a:solidFill"))
    if existing_fill is not None:
        spPr.remove(existing_fill)
    solidFill = etree.SubElement(spPr, qn("a:solidFill"))
    srgb = etree.SubElement(solidFill, qn("a:srgbClr"))
    srgb.set("val", _rgb_hex(color))
    alpha_elem = etree.SubElement(srgb, qn("a:alpha"))
    alpha_elem.set("val", str(int(alpha * 100000)))

    # subtle border
    ln = spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(spPr, qn("a:ln"))
    ln.set("w", "6350")  # 0.5pt
    for child in list(ln):
        ln.remove(child)
    if border_color:
        ln_fill = etree.SubElement(ln, qn("a:solidFill"))
        ln_srgb = etree.SubElement(ln_fill, qn("a:srgbClr"))
        ln_srgb.set("val", _rgb_hex(border_color))
        ln_alpha = etree.SubElement(ln_srgb, qn("a:alpha"))
        ln_alpha.set("val", str(int(border_alpha * 100000)))
    else:
        etree.SubElement(ln, qn("a:noFill"))

    # drop shadow
    add_card_shadow(shape, blur=76200, dist=25400, alpha=0.25)
    return shape


def add_gradient_line(slide, left, top, width, color1, color2, height=None):
    """그라데이션 강조선 (좌→우 페이드)"""
    if height is None:
        height = Pt(3)
    shape = slide.shapes.add_shape(1, left, top, width, height)
    sp = shape._element
    spPr = sp.spPr

    # Remove solid fill, add gradient
    existing = spPr.find(qn("a:solidFill"))
    if existing is not None:
        spPr.remove(existing)
    gradFill = etree.SubElement(spPr, qn("a:gradFill"))
    gsLst = etree.SubElement(gradFill, qn("a:gsLst"))
    gs1 = etree.SubElement(gsLst, qn("a:gs"))
    gs1.set("pos", "0")
    srgb1 = etree.SubElement(gs1, qn("a:srgbClr"))
    srgb1.set("val", _rgb_hex(color1))
    gs2 = etree.SubElement(gsLst, qn("a:gs"))
    gs2.set("pos", "100000")
    srgb2 = etree.SubElement(gs2, qn("a:srgbClr"))
    srgb2.set("val", _rgb_hex(color2))
    alpha_elem = etree.SubElement(srgb2, qn("a:alpha"))
    alpha_elem.set("val", "0")
    lin = etree.SubElement(gradFill, qn("a:lin"))
    lin.set("ang", "0")
    lin.set("scaled", "1")

    # no border
    ln = spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(spPr, qn("a:ln"))
    for child in list(ln):
        ln.remove(child)
    etree.SubElement(ln, qn("a:noFill"))
    return shape


def add_kpi_display(slide, left, top, value, label, value_color, label_color, palette,
                    value_size=48, label_size=12, width=None):
    """대형 KPI 숫자 + 라벨 (프로급 스타일)"""
    if width is None:
        width = Inches(3)
    # Value
    tb = add_text_box(slide, left, top, width, Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = value
    run.font.size = Pt(value_size)
    run.font.bold = True
    run.font.color.rgb = value_color
    run.font.name = palette.font
    # Label
    tb2 = add_text_box(slide, left, top + Inches(0.85), width, Inches(0.35))
    tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = label
    run2.font.size = Pt(label_size)
    run2.font.color.rgb = label_color
    run2.font.name = palette.font
    return tb, tb2


def add_deco_corner(slide, x, y, size, color, alpha=0.08):
    """데코 코너 삼각형 (장식 요소)"""
    shape = slide.shapes.add_shape(1, x, y, size, size)
    sp = shape._element
    spPr = sp.spPr
    prstGeom = spPr.find(qn("a:prstGeom"))
    if prstGeom is not None:
        prstGeom.set("prst", "rtTriangle")
    existing = spPr.find(qn("a:solidFill"))
    if existing is not None:
        spPr.remove(existing)
    solidFill = etree.SubElement(spPr, qn("a:solidFill"))
    srgb = etree.SubElement(solidFill, qn("a:srgbClr"))
    srgb.set("val", _rgb_hex(color))
    alpha_elem = etree.SubElement(srgb, qn("a:alpha"))
    alpha_elem.set("val", str(int(alpha * 100000)))
    ln = spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(spPr, qn("a:ln"))
    for child in list(ln):
        ln.remove(child)
    etree.SubElement(ln, qn("a:noFill"))
    return shape


# ---------------------------------------------------------------------------
# 레이아웃 빌더
# ---------------------------------------------------------------------------

def build_cover(slide, spec, palette):
    """1. cover — 히어로 표지 (v2.0 프로급)"""
    slide_w = Inches(13.333)
    slide_h = Inches(7.5)

    # 그라데이션 배경 (위→아래)
    set_gradient_bg(slide, palette.get("bg"), palette.get("bg_card"))

    # 데코 코너 (우상단)
    add_deco_corner(slide, Inches(11.5), Inches(-0.5), Inches(2.5), palette.get("accent1"), 0.06)
    add_deco_corner(slide, Inches(10.5), Inches(-1), Inches(3.5), palette.get("accent3"), 0.04)

    # 글로우 원 (더 풍부)
    add_glow_circle_v2(slide, Inches(-2), Inches(-1.5), Inches(6), palette.get("accent3"), 0.08)
    add_glow_circle_v2(slide, Inches(9.5), Inches(3.5), Inches(6), palette.get("accent1"), 0.07)
    add_glow_circle_v2(slide, Inches(4), Inches(5), Inches(4), palette.get("accent2"), 0.05)

    center_x = slide_w / 2
    margin_l = Inches(0.8)

    # 뱃지 (상단 중앙)
    badge_text = spec.get("badge", "")
    badge_color = palette.get("accent1")
    badge_w = Inches(3.5)
    if badge_text:
        add_badge(slide, center_x - badge_w / 2, Inches(1.2), badge_text, badge_color,
                  RGBColor(0xFF, 0xFF, 0xFF), width=badge_w)

    # 대형 타이틀
    title = spec.get("title", "")
    title_lines = title.split("\n") if title else [""]
    title_y = Inches(1.9)
    title_h = Inches(1.6)
    tb = add_text_box(slide, margin_l, title_y, slide_w - margin_l * 2, title_h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for line in title_lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        run.font.size = Pt(42)
        run.font.bold = True
        run.font.name = palette.font
        if title_lines.index(line) == 1:
            run.font.color.rgb = palette.get("accent2")
        else:
            run.font.color.rgb = palette.get("text")

    # 서브타이틀
    subtitle = spec.get("subtitle", "")
    if subtitle:
        tb2 = add_text_box(slide, margin_l, Inches(3.7), slide_w - margin_l * 2, Inches(0.5))
        tf2 = tb2.text_frame
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(16)
        run2.font.color.rgb = palette.get("text_secondary")
        run2.font.name = palette.font

    # 그라데이션 강조 라인 (좌→우 페이드아웃)
    line_w = Inches(3)
    add_gradient_line(slide, center_x - line_w / 2, Inches(4.35), line_w,
                      palette.get("accent1"), palette.get("accent2"), Pt(3))

    # body 텍스트 (회사 정보 등)
    body = spec.get("body", [])
    if body:
        y_body = Inches(4.7)
        for item in body:
            tb_b = add_text_box(slide, margin_l, y_body, slide_w - margin_l * 2, Inches(0.35))
            tf_b = tb_b.text_frame
            p_b = tf_b.paragraphs[0]
            p_b.alignment = PP_ALIGN.CENTER
            run_b = p_b.add_run()
            run_b.text = item
            run_b.font.size = Pt(13)
            run_b.font.color.rgb = palette.get("text_secondary")
            run_b.font.name = palette.font
            y_body += Inches(0.38)

    # 하단 장식선
    add_gradient_line(slide, Inches(0), Inches(7.1), slide_w,
                      palette.get("accent1"), palette.get("accent2"), Pt(2))


def build_three_cards(slide, spec, palette):
    """2. three_cards — 3-컬럼 카드 그리드 (v2.0 프로급)"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    cards = spec.get("cards", [])
    gap = Inches(0.3)
    card_w = Inches(3.7)
    card_h = Inches(4.5)
    slide_w = Inches(13.333)
    total_w = card_w * 3 + gap * 2
    start_x = (slide_w - total_w) / 2

    card_top = y + Inches(0.2)

    for i, card in enumerate(cards[:3]):
        cx = start_x + i * (card_w + gap)
        color_key = card.get("color", "accent1")
        accent = _resolve_color(color_key, palette)

        # 글래스 카드
        add_glass_card(slide, cx, card_top, card_w, card_h,
                       palette.get("bg_card"), alpha=0.3,
                       border_color=accent, border_alpha=0.15)

        # 상단 그라데이션 악센트
        add_gradient_line(slide, cx + Inches(0.2), card_top + Inches(0.1),
                          card_w - Inches(0.4), accent, palette.get("bg_card"), Pt(2))

        # 번호
        num_tb = add_text_box(slide, cx + Inches(0.25), card_top + Inches(0.3),
                               card_w - Inches(0.5), Inches(0.5))
        tf = num_tb.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = card.get("number", str(i + 1))
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = accent
        run.font.name = palette.font

        # 그라데이션 구분선
        add_gradient_line(slide, cx + Inches(0.25), card_top + Inches(0.95),
                          card_w - Inches(0.5), accent, palette.get("bg_card"), Pt(1))

        # 제목
        title_tb = add_text_box(slide, cx + Inches(0.25), card_top + Inches(1.1),
                                 card_w - Inches(0.5), Inches(0.45))
        tf2 = title_tb.text_frame
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        run2.text = card.get("title", "")
        run2.font.size = Pt(16)
        run2.font.bold = True
        run2.font.color.rgb = palette.get("text")
        run2.font.name = palette.font

        # 설명 (다중행)
        body = card.get("body", "")
        body_lines = body.split("\n") if body else []
        body_tb = add_text_box(slide, cx + Inches(0.25), card_top + Inches(1.65),
                                card_w - Inches(0.5), card_h - Inches(2.0))
        tf3 = body_tb.text_frame
        tf3.word_wrap = True
        first = True
        for line in body_lines:
            if first:
                p3 = tf3.paragraphs[0]
                first = False
            else:
                p3 = tf3.add_paragraph()
            p3.space_before = Pt(4)
            run3 = p3.add_run()
            run3.text = line
            run3.font.size = Pt(13)
            run3.font.color.rgb = palette.get("text_secondary")
            run3.font.name = palette.font


def build_two_column(slide, spec, palette):
    """3. two_column — 2-컬럼 레이아웃 (v2.0 프로급)"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    gap = Inches(0.3)
    card_w = Inches(5.7)
    card_h = Inches(5.5)
    card_top = y + Inches(0.15)

    # 좌측 글래스 카드
    add_glass_card(slide, margin_l, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=palette.get("accent1"), border_alpha=0.15)

    left_data = spec.get("left", {})
    left_title = left_data.get("title", "")
    if left_title:
        tb = add_text_box(slide, margin_l + Inches(0.3), card_top + Inches(0.3),
                           card_w - Inches(0.6), Inches(0.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = left_title
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = palette.get("text")
        run.font.name = palette.font

    items = left_data.get("items", [])
    item_y = card_top + Inches(0.85)
    for item in items:
        color_key = item.get("color", "accent1")
        accent = _resolve_color(color_key, palette)
        # 점
        add_icon_circle(slide, margin_l + Inches(0.3), item_y, "●", accent, Inches(0.3))
        tb2 = add_text_box(slide, margin_l + Inches(0.75), item_y, card_w - Inches(1.1), Inches(0.55))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        run2.text = item.get("year", "")
        run2.font.size = Pt(11)
        run2.font.bold = True
        run2.font.color.rgb = accent
        run2.font.name = palette.font

        p3 = tf2.add_paragraph()
        run3 = p3.add_run()
        run3.text = item.get("desc", "")
        run3.font.size = Pt(13)
        run3.font.color.rgb = palette.get("text")
        run3.font.name = palette.font
        item_y += Inches(0.85)

    # 우측 글래스 카드
    right_x = margin_l + card_w + gap
    add_glass_card(slide, right_x, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=palette.get("accent2"), border_alpha=0.15)

    right_data = spec.get("right", {})
    right_title = right_data.get("title", "")
    if right_title:
        tb = add_text_box(slide, right_x + Inches(0.3), card_top + Inches(0.3),
                           card_w - Inches(0.6), Inches(0.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = right_title
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = palette.get("text")
        run.font.name = palette.font

    stats = right_data.get("stats", [])
    stat_y = card_top + Inches(0.85)
    for stat in stats:
        color_key = stat.get("color", "accent1")
        accent = _resolve_color(color_key, palette)
        # 값
        tb_v = add_text_box(slide, right_x + Inches(0.3), stat_y, card_w - Inches(0.6), Inches(0.55))
        tf_v = tb_v.text_frame
        p_v = tf_v.paragraphs[0]
        run_v = p_v.add_run()
        run_v.text = stat.get("value", "")
        run_v.font.size = Pt(28)
        run_v.font.bold = True
        run_v.font.color.rgb = accent
        run_v.font.name = palette.font

        # 라벨 + 설명
        tb_l = add_text_box(slide, right_x + Inches(0.3), stat_y + Inches(0.55),
                             card_w - Inches(0.6), Inches(0.35))
        tf_l = tb_l.text_frame
        p_l = tf_l.paragraphs[0]
        run_l = p_l.add_run()
        run_l.text = f"{stat.get('label', '')}  {stat.get('desc', '')}"
        run_l.font.size = Pt(11)
        run_l.font.color.rgb = palette.get("text_secondary")
        run_l.font.name = palette.font

        # 프로그레스 바
        ratio = stat.get("ratio", 1.0)
        add_progress_bar(slide, right_x + Inches(0.3), stat_y + Inches(1.0),
                         card_w - Inches(0.6), ratio, accent)
        stat_y += Inches(1.5)


def build_three_layers(slide, spec, palette):
    """4. three_layers — 3-레이어 시스템 카드 (v2.0 프로급)"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    cards_data = spec.get("cards", [])
    gap = Inches(0.25)
    card_w = Inches(3.7)
    card_h = Inches(4.8)
    slide_w = Inches(13.333)
    total_w = card_w * 3 + gap * 2
    start_x = (slide_w - total_w) / 2
    card_top = y + Inches(0.2)

    arrow_colors = [palette.get("accent1"), palette.get("accent2")]

    for i, card in enumerate(cards_data[:3]):
        cx = start_x + i * (card_w + gap)
        color_key = card.get("color", "accent1")
        accent = _resolve_color(color_key, palette)

        # 글래스 카드
        add_glass_card(slide, cx, card_top, card_w, card_h,
                       palette.get("bg_card"), alpha=0.3,
                       border_color=accent, border_alpha=0.15)

        # 아이콘 원
        icon = card.get("icon", str(i + 1))
        add_icon_circle(slide, cx + Inches(0.3), card_top + Inches(0.3), icon, accent, Inches(0.6))

        # 레이어 번호
        layer_num = card.get("layer_num", f"Layer {i + 1}")
        tb_n = add_text_box(slide, cx + Inches(1.05), card_top + Inches(0.32),
                             card_w - Inches(1.2), Inches(0.3))
        tf_n = tb_n.text_frame
        p_n = tf_n.paragraphs[0]
        run_n = p_n.add_run()
        run_n.text = layer_num
        run_n.font.size = Pt(10)
        run_n.font.color.rgb = accent
        run_n.font.name = palette.font

        # 제목
        tb_t = add_text_box(slide, cx + Inches(0.25), card_top + Inches(1.05),
                             card_w - Inches(0.5), Inches(0.4))
        tf_t = tb_t.text_frame
        p_t = tf_t.paragraphs[0]
        run_t = p_t.add_run()
        run_t.text = card.get("title", "")
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = palette.get("text")
        run_t.font.name = palette.font

        # 설명
        desc = card.get("desc", "")
        tb_d = add_text_box(slide, cx + Inches(0.25), card_top + Inches(1.55),
                             card_w - Inches(0.5), Inches(0.6))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        run_d = p_d.add_run()
        run_d.text = desc
        run_d.font.size = Pt(12)
        run_d.font.color.rgb = palette.get("text_secondary")
        run_d.font.name = palette.font

        # 구분선
        add_accent_line(slide, cx + Inches(0.25), card_top + Inches(2.3),
                        card_w - Inches(0.5), palette.get("card_border"), Pt(1))

        # 기능 리스트
        features = card.get("features", [])
        feat_y = card_top + Inches(2.5)
        for feat in features[:4]:
            tb_f = add_text_box(slide, cx + Inches(0.35), feat_y,
                                 card_w - Inches(0.6), Inches(0.32))
            tf_f = tb_f.text_frame
            p_f = tf_f.paragraphs[0]
            run_f = p_f.add_run()
            run_f.text = f"· {feat}"
            run_f.font.size = Pt(11)
            run_f.font.color.rgb = palette.get("text_secondary")
            run_f.font.name = palette.font
            feat_y += Inches(0.35)

        # 카드 사이 화살표
        if i < 2:
            arrow_x = cx + card_w + gap / 2 - Inches(0.15)
            arrow_y = card_top + card_h / 2 - Inches(0.15)
            tb_a = add_text_box(slide, arrow_x, arrow_y, Inches(0.3), Inches(0.3))
            tf_a = tb_a.text_frame
            p_a = tf_a.paragraphs[0]
            p_a.alignment = PP_ALIGN.CENTER
            run_a = p_a.add_run()
            run_a.text = "→"
            run_a.font.size = Pt(18)
            run_a.font.color.rgb = arrow_colors[i % len(arrow_colors)]


def build_pipeline(slide, spec, palette):
    """5. pipeline — 4-단계 프로세스 플로우 (v2.0 프로급)"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    steps = spec.get("steps", [])
    kpis = spec.get("kpis", [])
    margin_l = Inches(0.8)
    gap = Inches(0.2)
    slide_w = Inches(13.333)
    n = max(len(steps), 1)
    total_gap = gap * (n - 1) + margin_l * 2
    card_w = (slide_w - total_gap) / n
    card_h = Inches(4.0)
    card_top = y + Inches(0.2)

    accent_keys = ["accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]

    for i, step in enumerate(steps):
        cx = margin_l + i * (card_w + gap)
        color_key = step.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)

        # 글래스 카드
        add_glass_card(slide, cx, card_top, card_w, card_h,
                       palette.get("bg_card"), alpha=0.3,
                       border_color=accent, border_alpha=0.15)
        # 상단 악센트 라인
        add_gradient_line(slide, cx + Inches(0.15), card_top + Inches(0.1),
                          card_w - Inches(0.3), accent, palette.get("bg_card"), Pt(2))

        # 번호 원
        num_size = Inches(0.55)
        add_icon_circle(slide, cx + Inches(0.3), card_top + Inches(0.3),
                        step.get("number", str(i + 1)), accent, num_size)

        # 제목
        tb_t = add_text_box(slide, cx + Inches(0.25), card_top + Inches(1.05),
                             card_w - Inches(0.5), Inches(0.45))
        tf_t = tb_t.text_frame
        p_t = tf_t.paragraphs[0]
        run_t = p_t.add_run()
        run_t.text = step.get("title", "")
        run_t.font.size = Pt(14)
        run_t.font.bold = True
        run_t.font.color.rgb = palette.get("text")
        run_t.font.name = palette.font

        # 설명
        desc = step.get("desc", "")
        tb_d = add_text_box(slide, cx + Inches(0.25), card_top + Inches(1.6),
                             card_w - Inches(0.5), card_h - Inches(1.9))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        run_d = p_d.add_run()
        run_d.text = desc
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = palette.get("text_secondary")
        run_d.font.name = palette.font

        # 화살표
        if i < len(steps) - 1:
            arrow_x = cx + card_w + gap / 2 - Inches(0.1)
            arrow_y = card_top + card_h / 2 - Inches(0.15)
            tb_a = add_text_box(slide, arrow_x, arrow_y, Inches(0.25), Inches(0.3))
            tf_a = tb_a.text_frame
            p_a = tf_a.paragraphs[0]
            p_a.alignment = PP_ALIGN.CENTER
            run_a = p_a.add_run()
            run_a.text = "→"
            run_a.font.size = Pt(14)
            run_a.font.color.rgb = palette.get("text_muted")

    # 하단 KPI 통계 바
    kpi_y = card_top + card_h + Inches(0.2)
    if kpis:
        kpi_w = (slide_w - margin_l * 2) / len(kpis)
        for i, kpi in enumerate(kpis):
            kx = margin_l + i * kpi_w
            color_key = kpi.get("color", accent_keys[i % len(accent_keys)])
            accent = _resolve_color(color_key, palette)
            tb_k = add_text_box(slide, kx, kpi_y, kpi_w - Inches(0.1), Inches(0.6))
            tf_k = tb_k.text_frame
            p_k = tf_k.paragraphs[0]
            run_k = p_k.add_run()
            run_k.text = kpi.get("value", "")
            run_k.font.size = Pt(20)
            run_k.font.bold = True
            run_k.font.color.rgb = accent
            run_k.font.name = palette.font
            p_k2 = tf_k.add_paragraph()
            run_k2 = p_k2.add_run()
            run_k2.text = kpi.get("label", "")
            run_k2.font.size = Pt(10)
            run_k2.font.color.rgb = palette.get("text_secondary")
            run_k2.font.name = palette.font


def build_split_detail(slide, spec, palette):
    """6. split_detail — 좌우 상세 분할"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    gap = Inches(0.3)
    slide_w = Inches(13.333)
    card_w = Inches(5.7)
    card_h = Inches(5.2)
    card_top = y + Inches(0.15)

    # 좌측: 항목 리스트 (글래스 카드)
    add_glass_card(slide, margin_l, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=palette.get("accent1"), border_alpha=0.15)

    items = spec.get("items", [])
    item_y = card_top + Inches(0.3)
    accent_keys = ["accent1", "accent2", "accent3", "accent4"]
    for i, item in enumerate(items):
        color_key = item.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)
        add_icon_circle(slide, margin_l + Inches(0.25), item_y, item.get("icon", str(i + 1)),
                        accent, Inches(0.5))
        tb = add_text_box(slide, margin_l + Inches(0.9), item_y, card_w - Inches(1.1), Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = item.get("name", "")
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = palette.get("text")
        run.font.name = palette.font
        p2 = tf.add_paragraph()
        run2 = p2.add_run()
        run2.text = item.get("desc", "")
        run2.font.size = Pt(11)
        run2.font.color.rgb = palette.get("text_secondary")
        run2.font.name = palette.font
        item_y += Inches(0.85)

    # 우측: 데이터 카드 (글래스)
    right_x = margin_l + card_w + gap
    add_glass_card(slide, right_x, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=palette.get("accent2"), border_alpha=0.15)

    data_items = spec.get("data", [])
    data_y = card_top + Inches(0.3)
    for i, di in enumerate(data_items):
        color_key = di.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)
        # 점수
        tb_s = add_text_box(slide, right_x + Inches(0.3), data_y, card_w - Inches(0.6), Inches(0.5))
        tf_s = tb_s.text_frame
        p_s = tf_s.paragraphs[0]
        run_s = p_s.add_run()
        run_s.text = f"{di.get('label', '')}  "
        run_s.font.size = Pt(13)
        run_s.font.color.rgb = palette.get("text_secondary")
        run_s.font.name = palette.font
        run_s2 = p_s.add_run()
        run_s2.text = di.get("score", "")
        run_s2.font.size = Pt(18)
        run_s2.font.bold = True
        run_s2.font.color.rgb = accent
        run_s2.font.name = palette.font
        add_progress_bar(slide, right_x + Inches(0.3), data_y + Inches(0.55),
                         card_w - Inches(0.6), di.get("ratio", 0.5), accent)
        data_y += Inches(1.0)


def build_execution(slide, spec, palette):
    """7. execution — 실행역량 + 예산"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    gap = Inches(0.3)
    slide_w = Inches(13.333)
    card_w = Inches(5.7)
    card_h = Inches(5.2)
    card_top = y + Inches(0.15)

    # 좌측: 역량 바 (글래스)
    add_glass_card(slide, margin_l, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=palette.get("accent1"), border_alpha=0.15)

    left_title = spec.get("left_title", "실행 역량")
    tb_lt = add_text_box(slide, margin_l + Inches(0.3), card_top + Inches(0.25),
                          card_w - Inches(0.6), Inches(0.38))
    tf_lt = tb_lt.text_frame
    p_lt = tf_lt.paragraphs[0]
    run_lt = p_lt.add_run()
    run_lt.text = left_title
    run_lt.font.size = Pt(14)
    run_lt.font.bold = True
    run_lt.font.color.rgb = palette.get("text")
    run_lt.font.name = palette.font

    skills = spec.get("skills", [])
    skill_y = card_top + Inches(0.75)
    accent_keys = ["accent1", "accent2", "accent3", "accent4", "accent5"]
    for i, sk in enumerate(skills):
        color_key = sk.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)
        tb_sk = add_text_box(slide, margin_l + Inches(0.3), skill_y, card_w - Inches(0.6), Inches(0.3))
        tf_sk = tb_sk.text_frame
        p_sk = tf_sk.paragraphs[0]
        run_sk = p_sk.add_run()
        run_sk.text = sk.get("name", "")
        run_sk.font.size = Pt(12)
        run_sk.font.bold = True
        run_sk.font.color.rgb = palette.get("text")
        run_sk.font.name = palette.font
        run_sk2 = p_sk.add_run()
        run_sk2.text = f"  {sk.get('desc', '')}"
        run_sk2.font.size = Pt(10)
        run_sk2.font.color.rgb = palette.get("text_muted")
        run_sk2.font.name = palette.font
        add_progress_bar(slide, margin_l + Inches(0.3), skill_y + Inches(0.35),
                         card_w - Inches(0.6), sk.get("ratio", 0.5), accent)
        skill_y += Inches(0.85)

    # 우측: 예산 breakdown (글래스)
    right_x = margin_l + card_w + gap
    add_glass_card(slide, right_x, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=palette.get("accent2"), border_alpha=0.15)

    right_title = spec.get("right_title", "예산 계획")
    tb_rt = add_text_box(slide, right_x + Inches(0.3), card_top + Inches(0.25),
                          card_w - Inches(0.6), Inches(0.38))
    tf_rt = tb_rt.text_frame
    p_rt = tf_rt.paragraphs[0]
    run_rt = p_rt.add_run()
    run_rt.text = right_title
    run_rt.font.size = Pt(14)
    run_rt.font.bold = True
    run_rt.font.color.rgb = palette.get("text")
    run_rt.font.name = palette.font

    budget_items = spec.get("budget", [])
    bud_y = card_top + Inches(0.75)
    for i, bi in enumerate(budget_items):
        color_key = bi.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)
        tb_bi = add_text_box(slide, right_x + Inches(0.3), bud_y, card_w - Inches(0.6), Inches(0.3))
        tf_bi = tb_bi.text_frame
        p_bi = tf_bi.paragraphs[0]
        run_bi = p_bi.add_run()
        run_bi.text = bi.get("name", "")
        run_bi.font.size = Pt(12)
        run_bi.font.color.rgb = palette.get("text")
        run_bi.font.name = palette.font
        run_bi2 = p_bi.add_run()
        pct = bi.get("ratio", 0.5)
        run_bi2.text = f"  {int(pct * 100)}%"
        run_bi2.font.size = Pt(11)
        run_bi2.font.bold = True
        run_bi2.font.color.rgb = accent
        run_bi2.font.name = palette.font
        add_progress_bar(slide, right_x + Inches(0.3), bud_y + Inches(0.35),
                         card_w - Inches(0.6), pct, accent)
        bud_y += Inches(0.85)


def build_revenue(slide, spec, palette):
    """8. revenue — 수익 모델 + 로드맵"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    slide_w = Inches(13.333)
    content_w = slide_w - margin_l * 2
    gap = Inches(0.3)

    # 상단: 3개 수익 카드
    revenue_cards = spec.get("revenue_cards", [])
    n = max(len(revenue_cards), 1)
    card_w = (content_w - gap * (n - 1)) / n
    card_h = Inches(2.8)
    card_top = y + Inches(0.15)

    accent_keys = ["accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]

    for i, rc in enumerate(revenue_cards):
        cx = margin_l + i * (card_w + gap)
        color_key = rc.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)
        add_glass_card(slide, cx, card_top, card_w, card_h,
                       palette.get("bg_card"), alpha=0.3,
                       border_color=accent, border_alpha=0.15)

        # 금액
        tb_v = add_text_box(slide, cx + Inches(0.25), card_top + Inches(0.25),
                             card_w - Inches(0.5), Inches(0.6))
        tf_v = tb_v.text_frame
        p_v = tf_v.paragraphs[0]
        run_v = p_v.add_run()
        run_v.text = rc.get("amount", "")
        run_v.font.size = Pt(26)
        run_v.font.bold = True
        run_v.font.color.rgb = accent
        run_v.font.name = palette.font

        # 제목
        tb_t = add_text_box(slide, cx + Inches(0.25), card_top + Inches(0.9),
                             card_w - Inches(0.5), Inches(0.35))
        tf_t = tb_t.text_frame
        p_t = tf_t.paragraphs[0]
        run_t = p_t.add_run()
        run_t.text = rc.get("title", "")
        run_t.font.size = Pt(13)
        run_t.font.bold = True
        run_t.font.color.rgb = palette.get("text")
        run_t.font.name = palette.font

        # 단위
        tb_u = add_text_box(slide, cx + Inches(0.25), card_top + Inches(1.3),
                             card_w - Inches(0.5), Inches(0.3))
        tf_u = tb_u.text_frame
        p_u = tf_u.paragraphs[0]
        run_u = p_u.add_run()
        run_u.text = rc.get("unit", "")
        run_u.font.size = Pt(10)
        run_u.font.color.rgb = palette.get("text_muted")
        run_u.font.name = palette.font

        # 설명
        tb_d = add_text_box(slide, cx + Inches(0.25), card_top + Inches(1.65),
                             card_w - Inches(0.5), card_h - Inches(1.9))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        run_d = p_d.add_run()
        run_d.text = rc.get("desc", "")
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = palette.get("text_secondary")
        run_d.font.name = palette.font

    # 합산 뱃지
    total = spec.get("total", "")
    if total:
        badge_w = Inches(3.5)
        add_badge(slide, slide_w / 2 - badge_w / 2, card_top + card_h + Inches(0.1),
                  total, palette.get("accent2"), RGBColor(0x0B, 0x0F, 0x1A), width=badge_w)

    # 하단: 타임라인
    timeline = spec.get("timeline", [])
    tl_y = card_top + card_h + Inches(0.6)
    tl_h = Inches(1.5)
    if timeline:
        # 선
        add_accent_line(slide, margin_l, tl_y + Inches(0.35), content_w,
                        palette.get("card_border"), Pt(2))
        step_w = content_w / max(len(timeline), 1)
        for i, tl in enumerate(timeline):
            tx = margin_l + i * step_w
            color_key = tl.get("color", accent_keys[i % len(accent_keys)])
            accent = _resolve_color(color_key, palette)
            # 점
            dot = slide.shapes.add_shape(1, tx + step_w / 2 - Inches(0.1),
                                          tl_y + Inches(0.25), Inches(0.2), Inches(0.2))
            dot.fill.solid()
            dot.fill.fore_color.rgb = accent
            dot.line.fill.background()
            sp_dot = dot._element
            prstGeom = sp_dot.spPr.find(qn("a:prstGeom"))
            if prstGeom is not None:
                prstGeom.set("prst", "ellipse")
            # 텍스트
            tb_tl = add_text_box(slide, tx, tl_y + Inches(0.6), step_w, Inches(0.8))
            tf_tl = tb_tl.text_frame
            tf_tl.word_wrap = True
            p_tl = tf_tl.paragraphs[0]
            p_tl.alignment = PP_ALIGN.CENTER
            run_tl = p_tl.add_run()
            run_tl.text = tl.get("year", "")
            run_tl.font.size = Pt(11)
            run_tl.font.bold = True
            run_tl.font.color.rgb = accent
            run_tl.font.name = palette.font
            p_tl2 = tf_tl.add_paragraph()
            p_tl2.alignment = PP_ALIGN.CENTER
            run_tl2 = p_tl2.add_run()
            run_tl2.text = tl.get("desc", "")
            run_tl2.font.size = Pt(10)
            run_tl2.font.color.rgb = palette.get("text_secondary")
            run_tl2.font.name = palette.font


def build_partners(slide, spec, palette):
    """9. partners — 파트너십 + KPI"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    gap = Inches(0.3)
    slide_w = Inches(13.333)
    card_w = Inches(5.7)
    card_h = Inches(5.2)
    card_top = y + Inches(0.15)

    # 좌측: 파트너 카드 리스트 (글래스)
    add_glass_card(slide, margin_l, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=palette.get("accent1"), border_alpha=0.15)

    partners = spec.get("partners", [])
    partner_y = card_top + Inches(0.25)
    accent_keys = ["accent1", "accent2", "accent3", "accent4", "accent5"]
    for i, partner in enumerate(partners):
        color_key = partner.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)
        # 이름
        tb_n = add_text_box(slide, margin_l + Inches(0.3), partner_y, card_w - Inches(0.6), Inches(0.35))
        tf_n = tb_n.text_frame
        p_n = tf_n.paragraphs[0]
        run_n = p_n.add_run()
        run_n.text = partner.get("name", "")
        run_n.font.size = Pt(14)
        run_n.font.bold = True
        run_n.font.color.rgb = accent
        run_n.font.name = palette.font
        # 역할
        tb_r = add_text_box(slide, margin_l + Inches(0.3), partner_y + Inches(0.38),
                             card_w - Inches(0.6), Inches(0.3))
        tf_r = tb_r.text_frame
        p_r = tf_r.paragraphs[0]
        run_r = p_r.add_run()
        run_r.text = partner.get("role", "")
        run_r.font.size = Pt(11)
        run_r.font.color.rgb = palette.get("text_secondary")
        run_r.font.name = palette.font
        # 태그
        tags = partner.get("tags", [])
        tag_x = margin_l + Inches(0.3)
        tag_y = partner_y + Inches(0.72)
        for tag in tags[:4]:
            add_badge(slide, tag_x, tag_y, tag, palette.get("bg_card_light"),
                      palette.get("text_muted"), width=Inches(1.3))
            tag_x += Inches(1.4)
        partner_y += Inches(1.35)

    # 우측: KPI 분해 (글래스)
    right_x = margin_l + card_w + gap
    add_glass_card(slide, right_x, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=palette.get("accent2"), border_alpha=0.15)

    kpi_title = spec.get("kpi_title", "KPI")
    tb_kt = add_text_box(slide, right_x + Inches(0.3), card_top + Inches(0.25),
                          card_w - Inches(0.6), Inches(0.38))
    tf_kt = tb_kt.text_frame
    p_kt = tf_kt.paragraphs[0]
    run_kt = p_kt.add_run()
    run_kt.text = kpi_title
    run_kt.font.size = Pt(14)
    run_kt.font.bold = True
    run_kt.font.color.rgb = palette.get("text")
    run_kt.font.name = palette.font

    kpis_q = spec.get("kpis_quantitative", [])
    kpis_nq = spec.get("kpis_nonquantitative", [])
    kpi_y = card_top + Inches(0.75)

    if kpis_q:
        tb_ql = add_text_box(slide, right_x + Inches(0.3), kpi_y, card_w - Inches(0.6), Inches(0.28))
        tf_ql = tb_ql.text_frame
        p_ql = tf_ql.paragraphs[0]
        run_ql = p_ql.add_run()
        run_ql.text = "계량 지표"
        run_ql.font.size = Pt(10)
        run_ql.font.bold = True
        run_ql.font.color.rgb = palette.get("accent2")
        run_ql.font.name = palette.font
        kpi_y += Inches(0.32)
        for kpi in kpis_q:
            color_key = kpi.get("color", "accent1")
            accent = _resolve_color(color_key, palette)
            tb_k = add_text_box(slide, right_x + Inches(0.3), kpi_y, card_w - Inches(0.6), Inches(0.28))
            tf_k = tb_k.text_frame
            p_k = tf_k.paragraphs[0]
            run_k = p_k.add_run()
            run_k.text = kpi.get("label", "")
            run_k.font.size = Pt(11)
            run_k.font.color.rgb = palette.get("text_secondary")
            run_k.font.name = palette.font
            run_k2 = p_k.add_run()
            run_k2.text = f"  {kpi.get('value', '')}"
            run_k2.font.size = Pt(12)
            run_k2.font.bold = True
            run_k2.font.color.rgb = accent
            run_k2.font.name = palette.font
            kpi_y += Inches(0.38)

    if kpis_nq:
        kpi_y += Inches(0.1)
        tb_nql = add_text_box(slide, right_x + Inches(0.3), kpi_y, card_w - Inches(0.6), Inches(0.28))
        tf_nql = tb_nql.text_frame
        p_nql = tf_nql.paragraphs[0]
        run_nql = p_nql.add_run()
        run_nql.text = "비계량 지표"
        run_nql.font.size = Pt(10)
        run_nql.font.bold = True
        run_nql.font.color.rgb = palette.get("accent4")
        run_nql.font.name = palette.font
        kpi_y += Inches(0.32)
        for kpi in kpis_nq:
            tb_k = add_text_box(slide, right_x + Inches(0.3), kpi_y, card_w - Inches(0.6), Inches(0.28))
            tf_k = tb_k.text_frame
            p_k = tf_k.paragraphs[0]
            run_k = p_k.add_run()
            run_k.text = f"· {kpi.get('label', '')}"
            run_k.font.size = Pt(11)
            run_k.font.color.rgb = palette.get("text_secondary")
            run_k.font.name = palette.font
            kpi_y += Inches(0.35)


def build_closing(slide, spec, palette):
    """10. closing — 클로징"""
    slide_w = Inches(13.333)
    slide_h = Inches(7.5)

    # 글로우 원
    add_glow_circle_v2(slide, Inches(-1.5), Inches(-1.5), Inches(6), palette.get("accent3"), 0.10)
    add_glow_circle_v2(slide, Inches(9.5), Inches(3.5), Inches(6), palette.get("accent1"), 0.08)

    margin_l = Inches(0.8)
    center_x = slide_w / 2

    # 큰 인용 기호
    tb_q = add_text_box(slide, margin_l, Inches(0.8), Inches(2), Inches(1.2))
    tf_q = tb_q.text_frame
    p_q = tf_q.paragraphs[0]
    p_q.alignment = PP_ALIGN.CENTER
    run_q = p_q.add_run()
    run_q.text = "\u201c"
    run_q.font.size = Pt(80)
    run_q.font.bold = True
    run_q.font.color.rgb = palette.get("accent2")
    run_q.font.name = palette.font

    # 대형 텍스트
    main_text = spec.get("main_text", spec.get("title", ""))
    lines = main_text.split("\n") if main_text else [""]
    tb_m = add_text_box(slide, margin_l, Inches(1.8), slide_w - margin_l * 2, Inches(2.0))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    first = True
    for i, line in enumerate(lines):
        if first:
            p = tf_m.paragraphs[0]
            first = False
        else:
            p = tf_m.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        run.font.size = Pt(36)
        run.font.bold = True
        run.font.name = palette.font
        if i == 1:
            run.font.color.rgb = palette.get("accent2")
        else:
            run.font.color.rgb = palette.get("text")

    # 강조 라인
    line_w = Inches(1.5)
    add_accent_line(slide, center_x - line_w / 2, Inches(4.0), line_w, palette.get("accent2"), Pt(3))

    # 회사 정보
    company = spec.get("company", "")
    if company:
        tb_c = add_text_box(slide, margin_l, Inches(4.5), slide_w - margin_l * 2, Inches(0.4))
        tf_c = tb_c.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.alignment = PP_ALIGN.CENTER
        run_c = p_c.add_run()
        run_c.text = company
        run_c.font.size = Pt(16)
        run_c.font.bold = True
        run_c.font.color.rgb = palette.get("text")
        run_c.font.name = palette.font

    # 연락처 카드
    contacts = spec.get("contacts", [])
    if contacts:
        contact_card_w = Inches(8)
        contact_card_h = Inches(1.0)
        cx = center_x - contact_card_w / 2
        add_glass_card(slide, cx, Inches(5.1), contact_card_w, contact_card_h,
                      palette.get("bg_card"), alpha=0.3, border_color=palette.get("accent1"), border_alpha=0.15)
        contact_text = "  |  ".join([f"{c.get('label', '')}: {c.get('value', '')}" for c in contacts])
        tb_ct = add_text_box(slide, cx + Inches(0.3), Inches(5.1) + Inches(0.3),
                              contact_card_w - Inches(0.6), Inches(0.4))
        tf_ct = tb_ct.text_frame
        p_ct = tf_ct.paragraphs[0]
        p_ct.alignment = PP_ALIGN.CENTER
        run_ct = p_ct.add_run()
        run_ct.text = contact_text
        run_ct.font.size = Pt(12)
        run_ct.font.color.rgb = palette.get("text_secondary")
        run_ct.font.name = palette.font


def build_key_statement(slide, spec, palette):
    """11. key_statement — 핵심 메시지 (v2.0 프로급)"""
    slide_w = Inches(13.333)
    slide_h = Inches(7.5)
    margin_l = Inches(0.8)

    # 데코 코너 (좌하단)
    add_deco_corner(slide, Inches(-0.5), Inches(5.5), Inches(3), palette.get("accent2"), 0.05)

    # 글로우 원 (2개로 풍부)
    add_glow_circle_v2(slide, slide_w / 2 - Inches(3), slide_h / 2 - Inches(3),
                       Inches(6), palette.get("accent2"), 0.08)
    add_glow_circle_v2(slide, Inches(8), Inches(-2), Inches(5), palette.get("accent1"), 0.05)

    # 핵심 메시지
    text = spec.get("title", spec.get("text", ""))
    lines = text.split("\n") if text else [""]
    tb = add_text_box(slide, margin_l, slide_h / 2 - Inches(1.2),
                      slide_w - margin_l * 2, Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for i, line in enumerate(lines):
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        run.font.size = Pt(38)
        run.font.bold = True
        run.font.name = palette.font
        if i == 1:
            run.font.color.rgb = palette.get("accent2")
        else:
            run.font.color.rgb = palette.get("text")

    # 그라데이션 강조선
    line_w = Inches(2)
    center_x = slide_w / 2
    add_gradient_line(slide, center_x - line_w / 2, slide_h / 2 + Inches(1.0), line_w,
                      palette.get("accent1"), palette.get("accent2"), Pt(3))

    # 서브타이틀
    subtitle = spec.get("subtitle", "")
    if subtitle:
        tb2 = add_text_box(slide, margin_l, slide_h / 2 + Inches(1.3),
                           slide_w - margin_l * 2, Inches(0.8))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(16)
        run2.font.color.rgb = palette.get("text_secondary")
        run2.font.name = palette.font

    # body 불릿
    body = spec.get("body", [])
    if body:
        # 글래스 카드 안에 배치
        card_w = Inches(10)
        card_x = center_x - card_w / 2
        card_y = slide_h / 2 + Inches(2.2)
        add_glass_card(slide, card_x, card_y, card_w, Inches(2.0),
                       palette.get("bg_card"), alpha=0.2,
                       border_color=palette.get("accent1"), border_alpha=0.15)
        body_y = card_y + Inches(0.2)
        for item in body:
            tb_b = add_text_box(slide, card_x + Inches(0.4), body_y, card_w - Inches(0.8), Inches(0.35))
            tf_b = tb_b.text_frame
            p_b = tf_b.paragraphs[0]
            p_b.alignment = PP_ALIGN.LEFT
            run_b = p_b.add_run()
            run_b.text = f"  {item}"
            run_b.font.size = Pt(13)
            run_b.font.color.rgb = palette.get("text_secondary")
            run_b.font.name = palette.font
            body_y += Inches(0.38)


def build_single_kpi(slide, spec, palette):
    """12. single_kpi — 단일 KPI 강조"""
    slide_w = Inches(13.333)
    slide_h = Inches(7.5)
    margin_l = Inches(0.8)

    y = _add_title_section(slide, spec, palette, Inches(0.5))

    color_key = spec.get("color", "accent1")
    accent = _resolve_color(color_key, palette)

    number = spec.get("number", "")
    label = spec.get("label", "")
    desc = spec.get("desc", "")
    ratio = spec.get("ratio", 0.75)

    center_x = slide_w / 2
    kpi_w = Inches(6)
    kpi_x = center_x - kpi_w / 2

    tb_n = add_text_box(slide, kpi_x, y + Inches(0.5), kpi_w, Inches(1.2))
    tf_n = tb_n.text_frame
    p_n = tf_n.paragraphs[0]
    p_n.alignment = PP_ALIGN.CENTER
    run_n = p_n.add_run()
    run_n.text = number
    run_n.font.size = Pt(72)
    run_n.font.bold = True
    run_n.font.color.rgb = accent
    run_n.font.name = palette.font

    tb_l = add_text_box(slide, kpi_x, y + Inches(1.8), kpi_w, Inches(0.45))
    tf_l = tb_l.text_frame
    p_l = tf_l.paragraphs[0]
    p_l.alignment = PP_ALIGN.CENTER
    run_l = p_l.add_run()
    run_l.text = label
    run_l.font.size = Pt(18)
    run_l.font.bold = True
    run_l.font.color.rgb = palette.get("text")
    run_l.font.name = palette.font

    tb_d = add_text_box(slide, kpi_x, y + Inches(2.35), kpi_w, Inches(0.4))
    tf_d = tb_d.text_frame
    p_d = tf_d.paragraphs[0]
    p_d.alignment = PP_ALIGN.CENTER
    run_d = p_d.add_run()
    run_d.text = desc
    run_d.font.size = Pt(13)
    run_d.font.color.rgb = palette.get("text_secondary")
    run_d.font.name = palette.font

    add_progress_bar(slide, kpi_x, y + Inches(2.9), kpi_w, ratio, accent, Pt(12))


def build_two_kpis(slide, spec, palette):
    """13. two_kpis — 2개 KPI 비교"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    slide_w = Inches(13.333)
    gap = Inches(0.3)
    card_w = Inches(5.7)
    card_h = Inches(4.5)
    card_top = y + Inches(0.2)

    kpis = spec.get("kpis", [{}, {}])
    accent_keys = ["accent1", "accent2"]

    for i, kpi in enumerate(kpis[:2]):
        cx = margin_l + i * (card_w + gap)
        color_key = kpi.get("color", accent_keys[i])
        accent = _resolve_color(color_key, palette)
        add_glass_card(slide, cx, card_top, card_w, card_h, palette.get("bg_card"),
                      alpha=0.3, border_color=accent, border_alpha=0.15)
        # 카드 상단 그라데이션 악센트 라인
        add_gradient_line(slide, cx, card_top, card_w, accent, palette.get("accent2"), Pt(3))
        tb_n = add_text_box(slide, cx + Inches(0.3), card_top + Inches(0.4),
                             card_w - Inches(0.6), Inches(1.0))
        tf_n = tb_n.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.alignment = PP_ALIGN.CENTER
        run_n = p_n.add_run()
        run_n.text = kpi.get("number", "")
        run_n.font.size = Pt(56)
        run_n.font.bold = True
        run_n.font.color.rgb = accent
        run_n.font.name = palette.font

        tb_l = add_text_box(slide, cx + Inches(0.3), card_top + Inches(1.55),
                             card_w - Inches(0.6), Inches(0.4))
        tf_l = tb_l.text_frame
        p_l = tf_l.paragraphs[0]
        p_l.alignment = PP_ALIGN.CENTER
        run_l = p_l.add_run()
        run_l.text = kpi.get("label", "")
        run_l.font.size = Pt(16)
        run_l.font.bold = True
        run_l.font.color.rgb = palette.get("text")
        run_l.font.name = palette.font

        tb_d = add_text_box(slide, cx + Inches(0.3), card_top + Inches(2.05),
                             card_w - Inches(0.6), Inches(0.4))
        tf_d = tb_d.text_frame
        p_d = tf_d.paragraphs[0]
        p_d.alignment = PP_ALIGN.CENTER
        run_d = p_d.add_run()
        run_d.text = kpi.get("desc", "")
        run_d.font.size = Pt(12)
        run_d.font.color.rgb = palette.get("text_secondary")
        run_d.font.name = palette.font

        ratio = kpi.get("ratio", 0.5)
        add_progress_bar(slide, cx + Inches(0.3), card_top + Inches(2.6),
                         card_w - Inches(0.6), ratio, accent)


def build_three_kpis(slide, spec, palette):
    """14. three_kpis — 3개 KPI 나란히 (v2.0 프로급)"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    slide_w = Inches(13.333)
    gap = Inches(0.3)
    card_w = Inches(3.7)
    card_h = Inches(4.5)
    total_w = card_w * 3 + gap * 2
    start_x = (slide_w - total_w) / 2
    card_top = y + Inches(0.2)

    kpis = spec.get("kpis", [{}, {}, {}])
    accent_keys = ["accent1", "accent2", "accent3"]

    for i, kpi in enumerate(kpis[:3]):
        cx = start_x + i * (card_w + gap)
        color_key = kpi.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)

        # 글래스 카드
        add_glass_card(slide, cx, card_top, card_w, card_h,
                       palette.get("bg_card"), alpha=0.3,
                       border_color=accent, border_alpha=0.2)

        # 상단 그라데이션 악센트 라인
        add_gradient_line(slide, cx + Inches(0.2), card_top + Inches(0.15),
                          card_w - Inches(0.4), accent, palette.get("bg_card"), Pt(2))

        # KPI 숫자
        add_kpi_display(slide, cx + Inches(0.1), card_top + Inches(0.5),
                        kpi.get("number", ""), kpi.get("label", ""),
                        accent, palette.get("text"), palette,
                        value_size=44, label_size=14, width=card_w - Inches(0.2))

        # 설명
        tb_d = add_text_box(slide, cx + Inches(0.25), card_top + Inches(2.0),
                             card_w - Inches(0.5), Inches(0.8))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.alignment = PP_ALIGN.CENTER
        run_d = p_d.add_run()
        run_d.text = kpi.get("desc", "")
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = palette.get("text_secondary")
        run_d.font.name = palette.font

        # 프로그레스 바
        ratio = kpi.get("ratio", 0.5)
        add_progress_bar(slide, cx + Inches(0.25), card_top + Inches(2.9),
                         card_w - Inches(0.5), ratio, accent)


def build_section_break(slide, spec, palette):
    """15. section_break — 섹션 구분"""
    slide_w = Inches(13.333)
    slide_h = Inches(7.5)

    add_glow_circle_v2(slide, slide_w / 2 - Inches(4), slide_h / 2 - Inches(4),
                       Inches(8), palette.get("accent3"), 0.08)

    margin_l = Inches(0.8)
    center_x = slide_w / 2

    section_num = spec.get("section_num", "")
    title = spec.get("title", "")
    subtitle = spec.get("subtitle", "")

    if section_num:
        add_icon_circle(slide, center_x - Inches(0.5), slide_h / 2 - Inches(2.0),
                        section_num, palette.get("accent1"), Inches(1.0))

    if title:
        tb = add_text_box(slide, margin_l, slide_h / 2 - Inches(0.8),
                          slide_w - margin_l * 2, Inches(0.9))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.size = Pt(42)
        run.font.bold = True
        run.font.color.rgb = palette.get("text")
        run.font.name = palette.font

    if subtitle:
        tb2 = add_text_box(slide, margin_l, slide_h / 2 + Inches(0.2),
                           slide_w - margin_l * 2, Inches(0.45))
        tf2 = tb2.text_frame
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(16)
        run2.font.color.rgb = palette.get("text_secondary")
        run2.font.name = palette.font


def build_comparison(slide, spec, palette):
    """16. comparison — 좌우 비교 (Before/After 또는 A vs B)"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    slide_w = Inches(13.333)
    card_w = Inches(5.5)
    card_h = Inches(5.0)
    gap = Inches(0.3)
    card_top = y + Inches(0.15)
    vs_w = Inches(0.9)

    left_data = spec.get("left", {})
    right_data = spec.get("right", {})
    vs_label = spec.get("vs", "VS")

    left_color = _resolve_color(left_data.get("color", "accent6"), palette)
    right_color = _resolve_color(right_data.get("color", "accent5"), palette)

    # 좌측 글래스 카드
    add_glass_card(slide, margin_l, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=left_color, border_alpha=0.2)
    add_gradient_line(slide, margin_l, card_top, card_w, left_color, palette.get("bg_card"), Pt(3))

    tb_lt = add_text_box(slide, margin_l + Inches(0.3), card_top + Inches(0.25),
                          card_w - Inches(0.6), Inches(0.45))
    tf_lt = tb_lt.text_frame
    p_lt = tf_lt.paragraphs[0]
    run_lt = p_lt.add_run()
    run_lt.text = left_data.get("title", "Before")
    run_lt.font.size = Pt(18)
    run_lt.font.bold = True
    run_lt.font.color.rgb = left_color
    run_lt.font.name = palette.font

    left_points = left_data.get("points", [])
    pt_y = card_top + Inches(0.85)
    for pt in left_points:
        tb_p = add_text_box(slide, margin_l + Inches(0.3), pt_y,
                             card_w - Inches(0.6), Inches(0.38))
        tf_p = tb_p.text_frame
        p_p = tf_p.paragraphs[0]
        run_p = p_p.add_run()
        run_p.text = f"· {pt}"
        run_p.font.size = Pt(13)
        run_p.font.color.rgb = palette.get("text_secondary")
        run_p.font.name = palette.font
        pt_y += Inches(0.42)

    # VS 구분자
    vs_x = margin_l + card_w + (slide_w - margin_l * 2 - card_w * 2) / 2 - vs_w / 2
    vs_y = card_top + card_h / 2 - Inches(0.3)
    tb_vs = add_text_box(slide, vs_x, vs_y, vs_w, Inches(0.6))
    tf_vs = tb_vs.text_frame
    p_vs = tf_vs.paragraphs[0]
    p_vs.alignment = PP_ALIGN.CENTER
    run_vs = p_vs.add_run()
    run_vs.text = vs_label
    run_vs.font.size = Pt(20)
    run_vs.font.bold = True
    run_vs.font.color.rgb = palette.get("text_muted")
    run_vs.font.name = palette.font

    # 우측 글래스 카드
    right_x = slide_w - margin_l - card_w
    add_glass_card(slide, right_x, card_top, card_w, card_h,
                   palette.get("bg_card"), alpha=0.3,
                   border_color=right_color, border_alpha=0.2)
    add_gradient_line(slide, right_x, card_top, card_w, right_color, palette.get("bg_card"), Pt(3))

    tb_rt = add_text_box(slide, right_x + Inches(0.3), card_top + Inches(0.25),
                          card_w - Inches(0.6), Inches(0.45))
    tf_rt = tb_rt.text_frame
    p_rt = tf_rt.paragraphs[0]
    run_rt = p_rt.add_run()
    run_rt.text = right_data.get("title", "After")
    run_rt.font.size = Pt(18)
    run_rt.font.bold = True
    run_rt.font.color.rgb = right_color
    run_rt.font.name = palette.font

    right_points = right_data.get("points", [])
    pt_y2 = card_top + Inches(0.85)
    for pt in right_points:
        tb_p2 = add_text_box(slide, right_x + Inches(0.3), pt_y2,
                              card_w - Inches(0.6), Inches(0.38))
        tf_p2 = tb_p2.text_frame
        p_p2 = tf_p2.paragraphs[0]
        run_p2 = p_p2.add_run()
        run_p2.text = f"· {pt}"
        run_p2.font.size = Pt(13)
        run_p2.font.color.rgb = palette.get("text_secondary")
        run_p2.font.name = palette.font
        pt_y2 += Inches(0.42)


def build_quote(slide, spec, palette):
    """17. quote — 인용구"""
    slide_w = Inches(13.333)
    slide_h = Inches(7.5)
    margin_l = Inches(1.2)

    add_glow_circle_v2(slide, Inches(0), slide_h / 2 - Inches(2.5),
                       Inches(5), palette.get("accent2"), 0.07)

    # 큰 따옴표
    tb_q = add_text_box(slide, margin_l, slide_h / 2 - Inches(2.2), Inches(1.5), Inches(1.5))
    tf_q = tb_q.text_frame
    p_q = tf_q.paragraphs[0]
    run_q = p_q.add_run()
    run_q.text = "\u201c"
    run_q.font.size = Pt(96)
    run_q.font.bold = True
    run_q.font.color.rgb = palette.get("accent2")
    run_q.font.name = palette.font

    # 인용 텍스트
    quote_text = spec.get("quote", spec.get("text", ""))
    tb_qt = add_text_box(slide, margin_l + Inches(0.5), slide_h / 2 - Inches(0.9),
                          slide_w - margin_l * 2, Inches(1.5))
    tf_qt = tb_qt.text_frame
    tf_qt.word_wrap = True
    p_qt = tf_qt.paragraphs[0]
    p_qt.alignment = PP_ALIGN.CENTER
    run_qt = p_qt.add_run()
    run_qt.text = quote_text
    run_qt.font.size = Pt(24)
    run_qt.font.italic = True
    run_qt.font.color.rgb = palette.get("text")
    run_qt.font.name = palette.font

    # 출처
    source = spec.get("source", "")
    if source:
        tb_s = add_text_box(slide, margin_l, slide_h / 2 + Inches(0.75),
                             slide_w - margin_l * 2, Inches(0.4))
        tf_s = tb_s.text_frame
        p_s = tf_s.paragraphs[0]
        p_s.alignment = PP_ALIGN.CENTER
        run_s = p_s.add_run()
        run_s.text = f"— {source}"
        run_s.font.size = Pt(13)
        run_s.font.color.rgb = palette.get("text_secondary")
        run_s.font.name = palette.font


def build_icon_row(slide, spec, palette):
    """18. icon_row — 아이콘 행"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    icons = spec.get("icons", [])
    slide_w = Inches(13.333)
    margin_l = Inches(0.8)
    n = max(len(icons), 1)
    item_w = (slide_w - margin_l * 2) / n
    icon_y = y + Inches(0.5)
    accent_keys = ["accent1", "accent2", "accent3", "accent4", "accent5"]

    for i, icon in enumerate(icons):
        ix = margin_l + i * item_w
        color_key = icon.get("color", accent_keys[i % len(accent_keys)])
        accent = _resolve_color(color_key, palette)
        icon_size = Inches(0.8)
        icon_x = ix + item_w / 2 - icon_size / 2
        add_icon_circle(slide, icon_x, icon_y, icon.get("icon", str(i + 1)), accent, icon_size)

        # 라벨
        tb_l = add_text_box(slide, ix, icon_y + icon_size + Inches(0.15), item_w, Inches(0.38))
        tf_l = tb_l.text_frame
        p_l = tf_l.paragraphs[0]
        p_l.alignment = PP_ALIGN.CENTER
        run_l = p_l.add_run()
        run_l.text = icon.get("label", "")
        run_l.font.size = Pt(13)
        run_l.font.bold = True
        run_l.font.color.rgb = palette.get("text")
        run_l.font.name = palette.font

        # 설명
        tb_d = add_text_box(slide, ix + Inches(0.1), icon_y + icon_size + Inches(0.6),
                             item_w - Inches(0.2), Inches(1.2))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.alignment = PP_ALIGN.CENTER
        run_d = p_d.add_run()
        run_d.text = icon.get("desc", "")
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = palette.get("text_secondary")
        run_d.font.name = palette.font


def build_data_table(slide, spec, palette):
    """19. data_table — 데이터 테이블"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    slide_w = Inches(13.333)
    content_w = slide_w - margin_l * 2

    headers = spec.get("headers", [])
    rows = spec.get("rows", [])

    if not headers and not rows:
        return

    n_cols = max(len(headers), max((len(r) for r in rows), default=1))
    col_w = content_w / n_cols
    row_h = Inches(0.45)
    header_h = Inches(0.48)
    table_top = y + Inches(0.15)

    # 헤더 행
    for j, header in enumerate(headers):
        hx = margin_l + j * col_w
        h_shape = slide.shapes.add_shape(1, hx, table_top, col_w - Inches(0.02), header_h)
        h_shape.fill.solid()
        h_shape.fill.fore_color.rgb = palette.get("accent1")
        h_shape.line.fill.background()
        tf_h = h_shape.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.alignment = PP_ALIGN.CENTER
        run_h = p_h.add_run()
        run_h.text = str(header)
        run_h.font.size = Pt(12)
        run_h.font.bold = True
        run_h.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run_h.font.name = palette.font

    # 데이터 행
    for ri, row in enumerate(rows):
        row_y = table_top + header_h + ri * row_h
        row_color = palette.get("bg_card") if ri % 2 == 0 else palette.get("bg_card_light")
        for ci, cell in enumerate(row[:n_cols]):
            cx = margin_l + ci * col_w
            c_shape = slide.shapes.add_shape(1, cx, row_y, col_w - Inches(0.02), row_h - Inches(0.02))
            c_shape.fill.solid()
            c_shape.fill.fore_color.rgb = row_color
            c_shape.line.fill.background()
            tf_c = c_shape.text_frame
            p_c = tf_c.paragraphs[0]
            p_c.alignment = PP_ALIGN.CENTER
            run_c = p_c.add_run()
            run_c.text = str(cell)
            run_c.font.size = Pt(11)
            run_c.font.color.rgb = palette.get("text_secondary")
            run_c.font.name = palette.font


def build_image_placeholder(slide, spec, palette):
    """20. image_placeholder — 이미지 자리 표시"""
    y = _add_title_section(slide, spec, palette, Inches(0.5))

    margin_l = Inches(0.8)
    slide_w = Inches(13.333)
    content_w = slide_w - margin_l * 2
    ph_h = Inches(4.5)

    # 회색 박스
    ph = slide.shapes.add_shape(1, margin_l, y + Inches(0.15), content_w, ph_h)
    ph.fill.solid()
    ph.fill.fore_color.rgb = RGBColor(0x25, 0x30, 0x48)
    ph.line.color.rgb = palette.get("card_border")
    ph.line.width = Pt(1)

    # 중앙 안내 텍스트
    tb = add_text_box(slide, margin_l, y + Inches(0.15) + ph_h / 2 - Inches(0.35),
                      content_w, Inches(0.7))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = spec.get("placeholder_text", "[이미지 영역]")
    run.font.size = Pt(18)
    run.font.color.rgb = palette.get("text_muted")
    run.font.name = palette.font

    # 캡션
    caption = spec.get("caption", "")
    if caption:
        tb_c = add_text_box(slide, margin_l, y + Inches(0.15) + ph_h + Inches(0.1),
                             content_w, Inches(0.35))
        tf_c = tb_c.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.alignment = PP_ALIGN.CENTER
        run_c = p_c.add_run()
        run_c.text = caption
        run_c.font.size = Pt(11)
        run_c.font.color.rgb = palette.get("text_muted")
        run_c.font.name = palette.font


# ---------------------------------------------------------------------------
# 레이아웃 맵
# ---------------------------------------------------------------------------

LAYOUT_MAP = {
    "cover": build_cover,
    "three_cards": build_three_cards,
    "two_column": build_two_column,
    "three_layers": build_three_layers,
    "pipeline": build_pipeline,
    "split_detail": build_split_detail,
    "execution": build_execution,
    "revenue": build_revenue,
    "partners": build_partners,
    "closing": build_closing,
    "key_statement": build_key_statement,
    "single_kpi": build_single_kpi,
    "two_kpis": build_two_kpis,
    "three_kpis": build_three_kpis,
    "section_break": build_section_break,
    "comparison": build_comparison,
    "quote": build_quote,
    "icon_row": build_icon_row,
    "data_table": build_data_table,
    "image_placeholder": build_image_placeholder,
}


# ---------------------------------------------------------------------------
# 메인 빌더
# ---------------------------------------------------------------------------

def build_presentation(plan, output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    palette = ColorPalette(
        plan.get("global_colors", DEFAULT_COLORS),
        plan.get("global_font", "Apple SD Gothic Neo"),
    )

    for slide_spec in plan.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
        # 프로급: 모든 슬라이드에 그라데이션 배경 적용
        set_gradient_bg(slide, palette.get("bg"), palette.get("bg_card"))
        layout_fn = LAYOUT_MAP.get(slide_spec.get("layout_id", ""), build_key_statement)
        layout_fn(slide, slide_spec, palette)
        if slide_spec.get("speaker_notes"):
            add_notes(slide, slide_spec["speaker_notes"])

    prs.save(output_path)
    return len(plan.get("slides", [])), os.path.getsize(output_path)


# ---------------------------------------------------------------------------
# CLI 진입점
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="python-pptx 슬라이드 빌더")
    parser.add_argument("--plan", required=True, help="slide_plan JSON 파일 경로")
    parser.add_argument("--output", required=True, help="출력 .pptx 파일 경로")
    args = parser.parse_args()

    with open(args.plan, "r", encoding="utf-8") as f:
        plan = json.load(f)

    slide_count, file_size = build_presentation(plan, args.output)
    file_size_kb = round(file_size / 1024, 1)

    result = {
        "success": True,
        "output_path": os.path.abspath(args.output),
        "slide_count": slide_count,
        "file_size_kb": file_size_kb,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
